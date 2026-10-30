"""阶段23证据门禁与白名单归档；默认只检查，缺少任何终态必须拒绝。"""
from pathlib import Path
import argparse
from collections import Counter
import copy
import datetime as dt
import gzip
import hashlib
import io
import json
import os
import re
import runpy
import shutil
import subprocess
import sys
import zipfile

ROOT = Path(__file__).resolve().parents[2]
WORK = Path(__file__).resolve().parent
BASE = 'd2dab71909d4c389482be9b95d9594db76310eeb'
OUT = ROOT / 'outputs/warehouses-20261030'
FORMAL = ROOT / 'docs/双仓储与五角色并发阶段验收.json'
FIXTURE = 'config/development/first-loop.json'
PUBLIC = 'client/assets/resources/development/gameplay.json'
GAME = 'client/assets/scenes/Game.scene'
ATTR_BASE = 'd278504d449285e20319dfb09dbc1d8e0ca6c7a9ec414efa8b475154af285bec'
ATTR_APPEND = b'outputs/warehouses-20261030/** -text\n'
REQUIRED_TESTS = {'TestWarehouseTransactionsSplitsLegacyReceiptsAndFencesMongo',
                  'TestWarehousesTwoRolesAtomicTransferAndReconnectMongoWS',
                  'TestWarehouseConcurrentFifthRoleAndSixthRejectedMongoHTTP',
                  'TestWarehousePermanentOriginsShareRunButNotAssetsMongo'}
CLIENT_COMMANDS = {'models', 'full-ui-app', 'scene', 'strict', 'protocol-check', 'protocol-lint'}
MONGO_KEYS = ('backend', 'pid', 'installation_id', 'project_key', 'data_path', 'uri', 'port', 'replica_set', 'version')
GUI_CONTRACT_REVIEWED = True  # 已静读冻结producer合同；实际两组终态/图片仍分别硬门禁，绝不预填通过。
GUI_TOOL_SHA = 'b34c99429aaa7834552f6ab7b19604480b72080a77c9acf453fd27b995c116b3'
GUI_CONTRACT_SHA = '183852eb6578ec2779b933ee9663b064f9c0b142aa0933039d7971fe57eed540'
QA_RUNNER_SHA = '2f657a4d95e4d743040dd9f3bafafa283e7e8539107f1ca3b708c27d5b8f3906'
CUTOVER_PEER_SHA = '31ac3e5da97300a139fbb1e0e8148d7023532708de639d7e1f04a77be789d330'


def need(ok, message):
    if not ok:
        raise ValueError(message)


def sha(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1048576), b''):
            digest.update(block)
    return digest.hexdigest()


def digest(raw):
    return hashlib.sha256(raw).hexdigest()


def local(value):
    raw = Path(value)
    raw = raw if raw.is_absolute() else ROOT / raw
    for parent in [raw, *raw.parents]:
        need(not parent.is_symlink() and not parent.is_junction(), '路径不能经过重解析点：' + str(parent))
        if parent == ROOT:
            break
    path = raw.resolve()
    need(path.is_relative_to(ROOT), '路径越界：' + str(value))
    return path


def relative(path):
    return local(path).relative_to(ROOT).as_posix()


def read(path):
    return json.loads(local(path).read_text(encoding='utf-8'))


def git(*args):
    return subprocess.check_output(['git', *args], cwd=ROOT, env={**os.environ, 'GIT_OPTIONAL_LOCKS': '0'})


def current(values, label):
    need(isinstance(values, dict) and values, label + '缺少输入集合')
    for name, expected in values.items():
        need(re.fullmatch(r'[0-9a-f]{64}', expected) and sha(local(name)) == expected, label + '当前原字节不同：' + name)


def snapshot(group):
    module = runpy.run_path(str(ROOT / 'tools/server/test.py' if group == 'go' else WORK / 'stage_checks.py'), run_name='readonly_enumeration')
    files = module['input_files' if group == 'go' else 'client_inputs']()
    return {relative(path): sha(path) for path in sorted(set(files))}


def tree(path):
    path = local(path)
    paths = list(path.rglob('*'))
    need(all(not p.is_symlink() and not p.is_junction() for p in paths), '产物目录包含链接')
    return {relative(p): sha(p) for p in sorted(paths) if p.is_file()}


def json_output(path):
    raw = local(path).read_text(encoding='utf-8')
    for text in [raw, *reversed(raw.splitlines())]:
        try:
            value = json.loads(text)
            if isinstance(value, dict):
                return value
        except json.JSONDecodeError:
            pass
    raise ValueError('缺JSON原始命令终态：' + str(path))


def report_log(name, report_path):
    p = Path(name)
    if p.is_absolute() or p.parts[0] in {'.work', 'client', 'server', 'tools', 'config', 'conf', 'docs'}:
        return local(p)
    need(len(p.parts) == 1, '未知支持日志相对路径接口')
    return local(report_path.parent / p)


def command_logs(report, passed=True):
    result = {}
    for row in report['commands']:
        name = row['name']
        need(name not in result and (not passed or row['exit_code'] == 0), '命令重复或非0终态：' + name)
        for stream in ['stdout', 'stderr']:
            need(sha(local(row[stream])) == row[stream + '_sha256'], '命令原始日志SHA不同：' + row[stream])
        result[name] = row
    return result


def external(report):
    module = runpy.run_path(str(WORK / 'stage_checks.py'), run_name='readonly_external_enumeration')
    for prefix, external_paths in [('external_engine', [module['ENGINE'] / name for name in module['ENGINE_FILES']]), ('external_font', module['FONTS'])]:
        values = report[prefix + '_before']
        need(values == report[prefix + '_after'] and set(values) == {str(p) for p in external_paths}, prefix + '前后/精确白名单不同')
        need(all(sha(Path(name)) == value for name, value in values.items()), prefix + '当前原字节不同')


def wrapper(path, group):
    value = read(path)
    need(value['group'] == group and value['passed'] is True and value['stable'] is True and value['freeze_started_before_commands'] is True, group + '包装终态未通过')
    need(value['before'] == value['after'] == snapshot(group) and value['input_count'] == len(value['before']), group + '输入前后/current或完整集合不同')
    commands = command_logs(value)
    need(set(commands) == (CLIENT_COMMANDS if group == 'client' else {'creator'} if group == 'creator' else {'full-mongo-race'}), group + '完整命令组不同')
    if group != 'go':
        external(value)
    return value, commands


def node_counts(path):
    values = {key: int(n) for key, n in re.findall(r'^(?:#|ℹ)\s+(tests|pass|fail|skipped|cancelled|todo)\s+(\d+)\s*$', local(path).read_text(encoding='utf-8'), re.M)}
    need(set(values) == {'tests', 'pass', 'fail', 'skipped', 'cancelled', 'todo'} and values['tests'] == values['pass'] > 0 and all(values[k] == 0 for k in ['fail', 'skipped', 'cancelled', 'todo']), 'Node原始计数没有完整通过：' + str(path))
    return values['pass']


def go_evidence(path):
    value = read(path)
    wrapped, _ = wrapper(path.parent / 'wrapper.json', 'go')
    need(value['passed'] is True and value['exit_code'] == 0 and value['mongo_e2e_explicit'] is True and value['race_detector'] is True, '全Go缺真正Mongo/race成功')
    need(value['command'] == [sys.executable, str(ROOT / 'tools/dev'), 'go', 'test', '-json', '-race', '-p', '1', './...', '-count=1', '-timeout=1200s'], 'Go命令必须为完整./.../race/p1/count1/1200，无过滤')
    need(value['package_parallelism'] == 1 and value['package_timeout_seconds'] == 1200 and value['single_rpc_and_business_deadlines_changed'] is False, '包预算与业务时限合同不符')
    need(value['source_changed_during_test'] is False and value['source_sha256'] == value['source_sha256_after'] == wrapped['after'] and len(value['source_sha256']) == 206, 'Go206输入不完整/漂移')
    raw = path.parent / 'events.jsonl'
    need(sha(raw) == value['events_sha256'], 'Go原始events未配对')
    events = [json.loads(line) for line in raw.read_text(encoding='utf-8').splitlines() if line.strip()]
    passed = [{'package': e['Package'], 'test': e['Test']} for e in events if e.get('Test') and e['Action'] == 'pass']
    skipped = [{'package': e['Package'], 'test': e['Test']} for e in events if e.get('Test') and e['Action'] == 'skip']
    failed = [{'package': e['Package'], 'test': e.get('Test')} for e in events if e['Action'] in {'fail', 'build-fail'}]
    need(passed == value['passed_tests'] and skipped == value['skipped_tests'] == [] and failed == value['failed_tests'] == [], 'Go raw统计不同或存在失败/跳过')
    need(REQUIRED_TESTS <= {e['test'] for e in passed if e['package'] == 'minimir/internal/integration'}, '缺4项真实仓储事务/双角色/抢5拒6/跨origin用例')
    need(value['fixture']['sha256'] == sha(ROOT / FIXTURE), 'Go不是当前fixture')
    return value, {'passed': len(passed), 'skipped': 0, 'failed': 0, 'inputs': len(value['source_sha256']), 'required_integration_tests': sorted(REQUIRED_TESTS)}


def expected_client_commands(path):
    return {
        'models': ['node', 'client/tests/ui/compile_models.mjs'],
        'full-ui-app': ['node', '--test', 'client/tests/ui/*.test.cjs', 'client/tests/app/*.test.cjs'],
        'scene': ['node', '--test', 'client/tests/scene/*.test.cjs'],
        'strict': ['node', 'tools/client/typecheck.mjs'],
        'protocol-check': ['task', 'protocol-check', '--report-dir', str(path.parent.relative_to(ROOT) / 'protocol')],
        'protocol-lint': ['python3', 'tools/config/generate_protocol.py', '--check']}


def client_evidence(path):
    value, commands = wrapper(path, 'client')
    app, scene = [node_counts(local(commands[n]['stdout'])) for n in ['full-ui-app', 'scene']]
    expected_commands = expected_client_commands(path)
    need(all(commands[n]['command'] == [sys.executable, str(ROOT / 'tools/dev'), *expected_commands[n]] for n in CLIENT_COMMANDS), 'client六命令被子集/过滤替代')
    tasks = [json.loads(line) for line in local(commands['protocol-check']['stdout']).read_text(encoding='utf-8').splitlines() if line.strip()]
    need(len(tasks) == 2 and all(row['exit_code'] == 0 for row in tasks), 'PB/绑定子命令缺失或失败')
    codec = [row for row in tasks if row['command'] == ['node', 'tools/protocol/verify_client_pb.cjs']]
    need(len(codec) == 1, '缺唯一原PB codec子报告')
    pb, lint = json_output(codec[0]['output']), json_output(commands['protocol-lint']['stdout'])
    need(pb['status'] == 'passed' and lint['valid'] is True and (pb['message_count'], pb['int64_field_count'], pb['int64_boundary_cases']) == (82, 195, 644), '协议静态codec原终态不同')
    need(pb['scope'] == 'candidate_static_protobuf_wire_codec_only' and pb['creator_project_build_verified'] is False, '不得把codec升格业务验收')
    return value, {'app': app, 'scene': scene, 'total': app + scene, 'inputs': value['input_count'], 'engine_inputs': len(value['external_engine_before']), 'font_inputs': len(value['external_font_before']),
                   'pb_messages': pb['message_count'], 'pb_long_fields': pb['int64_field_count'], 'pb_boundary_cases': pb['int64_boundary_cases'],
                   'pb_protocol_ids': lint['messages'], 'pb_error_codes': lint['error_codes']}, [local(t['output']) for t in tasks]


def protocol_delta():
    before = json.loads(git('show', BASE + ':shared/proto/protocol_manifest.json'))
    after = read(ROOT / 'shared/proto/protocol_manifest.json')
    for key, added in [('messages', {40020, 40021, 40022, 40023}), ('error_codes', {3501, 3502, 3503, 3504})]:
        old = {r['value']: r for r in before[key]}; new = {r['value']: r for r in after[key]}
        need(set(new) - set(old) == added and all(new.get(k) == v for k, v in old.items()), '协议旧定义变化或新增ID不精确')
    return {'new_message_ids': [40020, 40021, 40022, 40023], 'new_error_codes': [3501, 3502, 3503, 3504],
            'scope': 'codec静态与类型合同；业务另由真实Go/GUI证据，不将此映射视为业务通过'}


def configuration_evidence():
    review, run = read(WORK / 'configuration-review.json'), read(WORK / 'configuration-run-01/report.json')
    need(review['status'] == 'reviewed' and review['applied'] is True and review['head'] == BASE and run['passed'] is True, '配置作者审证/实际运行未通过')
    need(run['input_source_sha256'] == run['input_source_sha256_after'] and len(run['input_source_sha256']) == 27, '配置27源前后不同')
    current(run['input_source_sha256'], '配置27输入')
    current(run['generated_sha256'], '配置派生物')
    commands = command_logs(run)
    need(set(commands) == {'export', 'sync-bindings', 'fixture', 'pet-candidates', 'summon-candidates', 'export-check', 'fixture-check', 'pet-check', 'summon-check', 'bindings-check', 'python-full'}, '配置11命令不完整')
    stderr = local(commands['python-full']['stderr']).read_text(encoding='utf-8')
    counts = re.findall(r'^Ran (\d+) tests? in .*$', stderr, re.M)
    need(counts == ['140'] and re.search(r'^OK\s*$', stderr, re.M) and run['python_tests'] == 140, 'Python140原日志未通过')
    gameplay = 'outputs/development-prep-20261026/玩法配置.xlsx'
    architecture = 'outputs/development-prep-20261026/架构配置.xlsx'
    old_raw, new_raw = git('show', BASE + ':' + gameplay), local(gameplay).read_bytes()
    need(digest(new_raw) == review['candidate_sha256'] and sha(local(architecture)) == digest(git('show', BASE + ':' + architecture)), '唯一源表/架构书原SHA不同')
    with zipfile.ZipFile(io.BytesIO(old_raw)) as left, zipfile.ZipFile(io.BytesIO(new_raw)) as right:
        need(left.namelist() == right.namelist(), 'Excel ZIP成员集合/顺序变化')
        changed = [n for n in left.namelist() if left.read(n) != right.read(n)]
        need(changed == review['changed_zip_members'] and len(left.namelist()) - len(changed) == review['unchanged_zip_members'] == 102, '非编辑ZIP成员不同')
    import openpyxl
    before, after = [openpyxl.load_workbook(io.BytesIO(raw), data_only=False) for raw in [old_raw, new_raw]]
    need(before.sheetnames == after.sheetnames, '源表sheet集合变化')
    changes, formulas = [], 0
    for name in before.sheetnames:
        l, r = before[name], after[name]
        need((l.max_row, l.max_column) == (r.max_row, r.max_column), '源表尺寸变化')
        for rows in zip(l.iter_rows(), r.iter_rows()):
            for x, y in zip(*rows):
                if x.data_type == 'f':
                    formulas += 1
                    need(x.value == y.value and y.data_type == 'f', '原公式变化')
                if x.value != y.value:
                    changes.append({'sheet': name, 'row': x.row, 'column': x.column, 'before': x.value, 'after': y.value})
    need(changes == review['changed_cells'] and len(changes) == 2 and formulas == review['formulas_preserved'] == 460, '仅2格/460公式合同不同')
    need([(r['sheet'], r['row'], r['column']) for r in changes] == [('开发首循环夹具', 4, 3), ('开发首循环夹具', 4, 5)], '不是明确profile参数与来源说明两格')
    old, new = [json.loads(changes[0][k]) for k in ['before', 'after']]
    expected = copy.deepcopy(old); expected.update(rule_version=23, warehouses={'enabled': True, 'query_page_limit': 30})
    need(new == expected and review['existing_warehouse_rules_preserved'] is True and review['architecture_workbook_preserved'] is True, '不能改变其它数值或confirmed仓库规则')
    fixture, public, candidate = read(ROOT / FIXTURE), read(ROOT / PUBLIC), read(ROOT / 'config/development/summon-candidates.json')
    need(fixture['rule_version'] == public['rule_version'] == 23 and len(fixture['skills']) == 22 and not fixture['casting_supply']['enabled'] and fixture['pet_runtime']['bootstrap_mode'] == 'legacy_untrained_pet', '主22技能/legacy/护符门禁变化')
    need(17 not in [s['id'] for s in fixture['skills']] and candidate['runtime_enabled'] is False and candidate['required_rule_version'] == 17, '17候选被开放/降版')
    need(public['warehouses'] == fixture['warehouses'], '公开/服务端仓储参数不是同源')
    preview = read(ROOT / 'config/releases/preview_current.json')
    need(preview['path'] == preview['version'] and preview['version'].startswith('dev-warehouses-v0.23-') and read(ROOT / 'conf/toolchain/config-binding-sync.json')['config_version'] == preview['version'], 'preview与绑定不是本轮')
    validation = read(ROOT / 'config/reports/配置校验报告.json')
    need(validation['structure_valid'] is True and validation['errors'] == [] and validation['runtime_ready'] is False and validation['release_blockers'], '研发结构或未发布边界不同')
    return {'python': 140, 'inputs': 27, 'changed_cells': 2, 'formulas': 460, 'unchanged_zip_members': 102,
            'preview': preview['version'], 'fixture_sha256': sha(ROOT / FIXTURE), 'public_sha256': sha(ROOT / PUBLIC)}


def layout_evidence(layouts):
    need(len(layouts) == 10 and GAME in layouts, '不是10份保存布局')
    for name, expected in layouts.items():
        original = git('show', BASE + ':' + name)
        need(sha(local(name)) == expected, '当前保存布局漂移')
        if name != GAME:
            need(digest(original) == expected, '9份旧布局变化：' + name)
            continue
        before, after = json.loads(original), read(ROOT / name)
        need((len(before), len(after)) == (1191, 1373), '保存Game不是原1191+182')
        for index, old in enumerate(before):
            current_object = copy.deepcopy(after[index])
            if index in {253, 395}:
                need(old['__type__'] == current_object['__type__'] == 'cc.Node' and old['_name'] == current_object['_name'] == {253: 'Panels', 395: 'Page0'}[index], '旧父节点身份变化')
                appended = current_object['_children'][len(old['_children']):]
                need(current_object['_children'][:len(old['_children'])] == old['_children'] and len(appended) == 1 and appended[0]['__id__'] >= 1191 and after[appended[0]['__id__']]['_name'] == 'Warehouse', '不是唯一追加Warehouse子引用')
                current_object['_children'] = old['_children']
            need(current_object == old, '原保存对象额外变化：' + str(index))
    return {'unchanged_against_base': 9, 'intentionally_changed': [GAME], 'original_objects': 1191, 'appended_objects': 182,
            'old_object_changes': [{'index': 253, 'field': '_children', 'append': 'Warehouse'}, {'index': 395, 'field': '_children', 'append': 'Warehouse'}]}


def creator_evidence(path):
    value = read(path)
    wrapped, _ = wrapper(path.parent / 'wrapper.json', 'creator')
    need(value['creator'] == '3.8.8' and value['platform'] == 'web-mobile' and value['exit_code'] == 36 and value['success'] is True, '真实Creator尚未终态通过36')
    need(value['scene_prefab_layout_preserved'] is True and value['source_stable_during_build'] is True and value['served_build_preserved'] is True, '构建期间源/布局/旧主包变化')
    need(value['asset_hashes'] == {n: h for n, h in wrapped['after'].items() if Path(n).suffix in {'.scene', '.prefab'}}, 'Creator10布局与包装不配对')
    need(value['script_hashes'] == {n: h for n, h in wrapped['after'].items() if n.startswith('client/assets/scripts/') and Path(n).suffix in {'.ts', '.js'}}, 'Creator脚本与包装不配对')
    layout = layout_evidence(value['asset_hashes'])
    before, after = [read(path.parent / n) for n in ['served-build-before.json', 'served-build-after.json']]
    need(before == after and len(before) == value['served_build_files'] > 0, '旧主包构建期间原清单不保留')
    root, manifest, config = [local(value[k]) for k in ['build_root', 'artifact_manifest', 'build_config']]
    need(root == path.parent / 'candidate/web-mobile' and manifest == path.parent / 'candidate-files.json', '候选目录/manifest身份错误')
    need(sha(manifest) == value['artifact_manifest_sha256'] and tree(root) == read(manifest) and len(read(manifest)) == value['artifact_file_count'] > 0, '真实候选产物SHA不同')
    need(sha(config) == value['build_config_sha256'] and sha(ROOT / 'conf/creator-build.json') == value['original_build_config_sha256'], '候选构建配置身份不同')
    expected = read(ROOT / 'conf/creator-build.json'); expected['buildPath'] = str(root.parent)
    need(expected == read(config), '候选构建配置有额外改变')
    need((path.parent / 'creator.log').is_file(), '缺原Creator log')
    return value, layout


def publication_evidence(path, creator, source_review, selected):
    transition = read(source_review)
    source = (WORK / 'cutover.py').read_bytes()
    original = source
    need(set(selected) == {'go', 'client', 'creator'} and transition['selected_attempts'] == selected, '切主review声明与CLI三类selector不同')
    for group, attempt in selected.items():
        need(re.fullmatch(r'[1-9][0-9]*', attempt), '非法selector编号')
        token = (group + '/full-' + attempt).encode()
        need(token in original, '切主source缺该精确selector：' + token.decode())
        original = original.replace(token, (group + '/full-1').encode())
    need(sha(WORK / 'cutover.py') == transition['current_source_sha256'] and digest(original) == transition['reviewed_source_sha256'] == CUTOVER_PEER_SHA and transition['all_other_source_bytes_equal'] is True, '切主peer旧源与新selector没有精确原字节桥接')
    report = read(path)
    need(report['passed'] is True and report['phase'] == 'completed' and report['rollback_attempted'] is False and report['rollback_restored'] is False and not report.get('failure') and not report.get('error'), '切主未真通过或有回滚')
    commands = command_logs(report)
    expected = ['mongo-before', 'server-before', 'client-before', 'server-stop', 'client-stop', 'client-stopped-status', 'server-start', 'client-start', 'client-after', 'mongo-after']
    need(list(commands) == expected, '切主必须精确10项正常命令，不能夹restart Mongo等')
    need(report['other_ports_touched'] == [] and report['mongo_service_restarted'] is False and report['served_candidate_exact_match'] is True, '切主触及非本任务服务或未核候选')
    mongo = [json_output(commands[key]['stdout']) for key in ['mongo-before', 'mongo-after']]
    need(all(all(v[k] is True for k in ['running', 'endpoint_verified', 'ownership_verified', 'listener_verified']) for v in mongo), 'Mongo原所有权不完整')
    need(all(mongo[0][k] == mongo[1][k] for k in MONGO_KEYS), 'Mongo九身份变化，不能称未重启')
    need(report['mongo_before'] == {k: mongo[0][k] for k in MONGO_KEYS} and report['mongo_after'] == {k: mongo[1][k] for k in MONGO_KEYS}, '切主Mongo摘要未配对')
    before, after = read(WORK / 'runtime-before.json'), read(WORK / 'runtime-after.json')
    need(report['before'] == before and report['after'] == after, '原runtime前后与切主报告不同')
    old_fixture = read(WORK / 'configuration-baseline.json')['files'][FIXTURE]
    for port, mode in [(18100, 'account'), (18101, 'game'), (18102, 'gm')]:
        old, new = before[str(port)], after[str(port)]
        need(old['mode'] == new['mode'] == mode and old['status'] == new['status'] == 'ok' and old['rule_version'] == 22 and old['fixture_hash'] == old_fixture, '切主前不是真实22原fixture')
        need(new['rule_version'] == 23 and new['fixture_hash'] == sha(ROOT / FIXTURE) and old['instance'] != new['instance'], '切主后不是真实23原fixture/新实例')
    manifest = read(local(creator['artifact_manifest'])); root = local(creator['build_root'])
    served = ROOT / 'client/build/web-mobile'
    expected_files = {(local(n).relative_to(root)).as_posix(): h for n, h in manifest.items()}
    actual = {p.relative_to(served).as_posix(): sha(p) for p in served.rglob('*') if p.is_file()}
    need(actual == expected_files and report['candidate_manifest_sha256'] == creator['artifact_manifest_sha256'], '主包不是原真实候选全量原字节')
    need(after['18080']['port'] == 18080 and after['18080']['account_port'] == 18100 and after['18080']['game_port'] == 18101 and Path(after['18080']['build_root']).resolve() == served and before['18080']['run_id'] != after['18080']['run_id'], '主静态身份/路由没有受控切换')
    return report


def gui_evidence(path, mode, creator_path, contract_path=None, contract_sha=None):
    need(GUI_CONTRACT_REVIEWED, 'PENDING：正常/隔离GUI实际raw schema尚待逐字段审阅，禁止提前验收')
    contract_path = contract_path or WORK / 'browser/report-contract-3.json'
    need(sha(contract_path) == (contract_sha or GUI_CONTRACT_SHA), 'GUI producer合同不是显式冻结版本')
    contract = read(contract_path)
    need(contract['tool'] == 'tools/client/verify_warehouses_browser.cjs', '不接受另一个相似GUI工具')
    static_path = local(contract['static_report']); static = read(static_path)
    need(sha(static_path) == contract['static_report_sha256'] and static['passed'] is True and static['before'] == static['after'] == static['current_frozen'], 'GUI工具真实静态报告不完整/前后不同')
    current(static['after'], 'GUI静态冻结当前输入')
    need(len(static['commands']) == 5, 'GUI静态完整命令组不是5项')
    for row in static['commands']:
        need(row['exit_code'] == 0 and '--run' not in row['command'], 'GUI静态命令不应执行浏览器/环境')
        for stream in ['stdout', 'stderr']:
            need(sha(report_log(row[stream], static_path)) == row[stream + '_sha256'], 'GUI静态原日志SHA不同')
        if '--test' in row['command']:
            need(node_counts(report_log(row['stdout'], static_path)) > 0, 'GUI工具纯测原终态缺失')
    value = read(path)
    need(set(contract['top_success_required']) <= value.keys() and value['passed'] is True and value['mode'] == mode and value['expectedRuleVersion'] == 23 and value['errors'] == [] and not value.get('error'), 'GUI未实际完整终态通过')
    need(value['toolSha256Before'] == value['toolSha256After'] == sha(ROOT / contract['tool']) == contract['tool_sha256'], 'GUI实际工具前后/current不是冻结source')
    need(value['fixtureSha256'] == value['fixtureSha256After'] == sha(ROOT / FIXTURE) and value['rawPublicSha256'] == value['rawPublicSha256After'] == sha(ROOT / PUBLIC), 'GUI fixture/public前后/current不同')
    need(local(value['creatorReportPath']) == creator_path and value['creatorReportSha256'] == value['creatorReportSha256After'] == sha(creator_path), 'GUI不是当前选定真实Creator')
    creator = read(creator_path)
    need(local(value['creatorBuildRoot']) == local(creator['build_root']) and value['sourceBefore'] == value['sourceAfter'] == {**creator['asset_hashes'], **creator['script_hashes']}, 'GUI输入集合不是原Creator源码/布局')
    current(value['sourceAfter'], 'GUI当前输入')
    need(value['touchInput'] is True and all(value[k] is False for k in ['routeOverride', 'privateBusinessInvoked', 'bodySeeded', 'naturalPaidChainVerified', 'physicalMobileVerified', 'productionModified']) and value['qualificationSeeded'] is (mode == 'isolated'), 'GUI不能使用私有业务、route override或假身体/真实付费宣称')
    need(value['healthBefore'] == value['healthAfter'] and len(value['healthBefore']) == 4, 'GUI前后4端点身份不保持')
    expected_ports = {18100: 'account', 18101: 'game', 18102: 'gm'} if mode == 'normal' else {18110: 'account', 18111: 'game', 18112: 'gm'}
    health = {row['port']: row for row in value['healthAfter']}
    static_port = 18080 if mode == 'normal' else 18082
    need(set(health) == set(expected_ports) | {static_port} and value['url'] == f'http://127.0.0.1:{static_port}/', 'GUI端点或原始URL不符')
    for port, expected_mode in expected_ports.items():
        row = health[port]
        need(row['status'] == 'ok' and row['mode'] == expected_mode and row['rule_version'] == 23 and row['fixture_hash'] == sha(ROOT / FIXTURE) and row['instance'], 'GUI三实例不是真实23/当前fixture')
    loaded = value['loadedPublic']; snapshot_path = path.parent / loaded['path']
    need(loaded['path'] == 'loaded-public.json' and sha(snapshot_path) == loaded['sha256'] and loaded['rawPublicSha256'] == sha(ROOT / PUBLIC) and loaded['nodeDeepStrictEqual'] is True and loaded['nodeVersion'] == value['nodeVersion'], '真实cc.resources快照/Node全对象证据缺失')
    need(json.loads(snapshot_path.read_text(encoding='utf-8')) == read(ROOT / PUBLIC), '实际loaded public全字段值与当前源不同')
    visual = read(path.parent / 'visual-review.json')
    need(visual['visual_passed'] is True and visual['report_sha256'] == sha(path) and visual['screenshots'] == value['screenshots'] and len(set(value['screenshots'])) == len(value['screenshots']) > 0, '缺全部实际图片复核')
    for name in value['screenshots']:
        need(Path(name).name == name and name.endswith('.png') and sha(path.parent / name) == visual['screenshot_sha256'][name] and (path.parent / name).read_bytes()[:8] == b'\x89PNG\r\n\x1a\n', '截图原SHA或PNG格式不符')
    observations = value['observations']
    need(set(contract['common_observations'] + contract[mode + '_observations']) <= observations.keys(), 'GUI缺完整业务观察截面')
    initial = observations['initial']
    battle, town = observations['battleEntry'], observations['townEntry']
    need(battle['map'] == '1' and town['map'] == initial['map'] == '2' and battle['scene'] != town['scene'] and battle['roleKey'] == town['roleKey'] == initial['roleKey'] and battle['life'] == town['life'] and int(town['hp']) > 0, '仓储GUI没有先经正常入口进入安全Town原birth')
    need(observations['townTravel']['before'] == battle and observations['townTravel']['after'] == town and observations['townTravel']['confirmation'], 'Town旅行没有真实确认前后截面')
    plans = [r for r in value['travelPlans'] if r['fromScene'] == battle['scene'] and r['map'] == '2' and r['binding'] == battle['binding']]
    need(plans and any(r['phase'] == 2 and r['id'] in {p['id'] for p in plans} and r['fromScene'] == battle['scene'] and r['scene'] == town['scene'] and r['map'] == '2' and r['binding'] == town['binding'] for r in value['travelStates']), 'Town新Entry没有20008/20010同旅行ID提交事实')
    need(len(value['createdRoles']) == 1 and value['createdRoles'][0]['job'] == initial['job'] == 2 and initial['roleKey'] == value['createdRoles'][0]['origin'] + ':' + value['createdRoles'][0]['roleId'], 'GUI不是正常HTTP创角的新道士')
    need(initial['rule'] == '23' and initial['auto'] is False and initial['hostingAuto'] is False, 'GUI初态没有正常关闭Auto')
    config = read(ROOT / FIXTURE)['warehouses']
    need(observations['personalInitial']['warehouse']['capacity'] == str(config['personal_initial_slots']) and observations['personalInitial']['warehouse']['items'] == [], '个人仓不是初始空仓容量')
    successful = [row for row in value['actions'] if not row['cancelled']]
    requests = [row for row in value['requests'] if row['id'] == 40022]
    need(len(successful) == len(requests) and len({row['opId'] for row in requests}) == len(requests), '40022与明确确认动作数量不同/重复发新op')
    for action in value['actions']:
        need(set(contract['action_fields']) <= action.keys(), '动作缺before/after/明确确认')
        before, after = action['before'], action['after']
        need(action['confirmation']['title'] and action['confirmation']['message'] and before['roleKey'] == after['roleKey'] == initial['roleKey'] and before['map'] == after['map'] == '2' and before['life'] == after['life'] and after['auto'] is False and after['hostingAuto'] is False, '动作角色/原birth/Town/用户Auto或确认截面不符')
        if action['cancelled']:
            need(before['items'] == after['items'] and before['wallet'] == after['wallet'] and before['warehouse'] == after['warehouse'] and not [q for q in requests if before['time'] <= q['time'] <= after['time']], '取消发出资产请求或改变原物')
            continue
        need(set(contract['committed_action_extra_fields']) <= action.keys(), '提交动作缺原协议事实')
        q, op, receipt = action['request'], action['operationResult'], action['actionResult']
        need(q in requests and op in value['operationResults'] and receipt in value['actionResults'] and re.fullmatch(r'[0-9a-f]{32}', q['opId']), '动作没有匹配实际raw帧')
        need(q['binding'] == before['binding'] == after['binding'] == op['binding'] == receipt['binding'] and q['request'] == op['request'] == receipt['request'], '原40022/40005/40023请求或binding错配')
        need(op['opId'] == receipt['opId'] == q['opId'] and op['pending'] is False and op['error'] == 0 and op['command'] == q['command'], '40005不是同op可信终态')
        need(receipt['role'] == before['roleKey'] and receipt['scope'] == q['scope'] == before['warehouse']['scope'] == after['warehouse']['scope'] and receipt['action'] == q['action'], '40023角色/scope/action不是原40022')
        need(q['rule'] == '23' and q['warehouseRevision'] == before['warehouse']['revision'] and int(receipt['warehouseRevision']) == int(q['warehouseRevision']) + 1 and after['warehouse']['revision'] == receipt['warehouseRevision'], '40023仓库版本不是原rev+1')
        validate_fresh_revision(q, op, receipt, before, after)
        need(before['time'] <= q['time'] <= after['time'] and op['time'] >= q['time'] and receipt['time'] >= q['time'], '原请求时间截面不符')
        need(after['warehouse']['gold'] == receipt['gold'] and after['warehouse']['yuanbaoUnits'] == receipt['yuanbaoUnits'] and after['warehouse']['capacity'] == receipt['capacity'], '重新查询页不等于已确认原仓库事实')
        need(any(r['binding'] == after['binding'] and r['revision'] == after['revision'] and r['items'] == after['items'] and r['equipped'] == after['equipped'] for r in value['inventories']), '动作after背包不是实际40003原截面')
        need(any(r['binding'] == after['binding'] and r['revision'] == after['revision'] and all(r[k] == after['wallet'][k] for k in ['gold', 'boundGold', 'yuanbaoUnits', 'boundYuanbaoUnits']) for r in value['wallets']), '动作after钱包不是实际40004原截面')
        need(any(r['binding'] == after['binding'] and r['role'] == after['roleKey'] and all(r[k] == after['warehouse'][k] for k in ['scope', 'revision', 'capacity', 'startSlot', 'limit', 'items', 'gold', 'yuanbaoUnits']) for r in value['pages']), '动作after仓页不是实际40021原截面')
        if q['action'] in (1, 2):
            validate_item_action(q, receipt, before, after)
        elif q['action'] in (3, 4):
            need(receipt['changes'] == [] and before['items'] == after['items'] and before['warehouse']['items'] == after['warehouse']['items'], '货币动作不应改变背包或当前仓页原件')
            key, sign = ('gold' if q['currency'] == 1 else 'yuanbaoUnits'), (-1 if q['action'] == 3 else 1)
            delta = positive(q['amount'])
            need(int(after['wallet'][key]) == int(before['wallet'][key]) + sign * delta and int(after['warehouse'][key]) == int(before['warehouse'][key]) - sign * delta, '原币种/整数单位不守恒')
            need(after['wallet']['boundGold'] == before['wallet']['boundGold'] and after['wallet']['boundYuanbaoUnits'] == before['wallet']['boundYuanbaoUnits'], '货币动作消耗了绑定钱包')
        else:
            need(receipt['changes'] == [] and before['items'] == after['items'] and before['warehouse']['items'] == after['warehouse']['items'], '扩页不应改变背包或当前仓页原件')
            need(q['action'] == 5 and q['scope'] == 1 and int(after['warehouse']['capacity']) == int(before['warehouse']['capacity']) + config['personal_expansion_slots'] and int(after['wallet']['yuanbaoUnits']) == int(before['wallet']['yuanbaoUnits']) - config['personal_expansion_yuanbao_units'], '个人扩页没有同源非绑元宝/容量守恒')
            need(after['wallet']['boundYuanbaoUnits'] == before['wallet']['boundYuanbaoUnits'] and '非绑定元宝' in action['confirmation']['message'], '扩页未明确非绑确认')
    for left, right in [('beforeReentry', 'reentry')] + ([('beforeFinalReentry', 'finalReentry')] if mode == 'isolated' else []):
        before, after = observations[left], observations[right]
        need(before['binding'] != after['binding'] and all(before[k] == after[k] for k in ['roleKey', 'scene', 'life', 'auto', 'hostingAuto', 'wallet']) and sorted(before['items'], key=lambda x: x['id']) == sorted(after['items'], key=lambda x: x['id']), '健康新票改变原场景/birth/资产/Auto')
        need(not [q for q in requests if before['time'] <= q['time'] <= after['time']] and not [r for r in value['actionResults'] if r['binding'] == after['binding'] and r['time'] <= after['time']], '健康新票重发旧40022或补播旧40023')
    for row in value['selections']:
        need(set(contract['selection_fields']) <= row.keys() and row['touch']['fixedButtonOverlap'] is False, '原条目触点没有真实保存Mask/按钮边界观察')
        need(row['after'][{'Inventory': 'selectedInventory', 'BagList': 'selectedBag', 'WarehouseList': 'selectedSlot'}[row['listName']]] == row['key'], '触屏选择没有落到原UID/仓位')
        for key in (['inventoryOffset'] if row['listName'] == 'Inventory' else ['bagOffset', 'warehouseOffset']):
            need(abs(row['beforeTouch'][key] - row['after'][key]) <= 1, '选取触点改变了原滚动位置')
        need(row['beforeTouch']['nav'] == row['after']['nav'], '面板选取清除了地面导航')
    details = observations['personalDetails']
    need(details['shown']['detail'] and all(abs(details['before'][key] - details['after'][key]) <= 1 for key in ['bagOffset', 'warehouseOffset']) and details['before']['nav'] == details['after']['nav'], '保存详情/返回原列表偏移或导航未保持')
    if mode == 'normal':
        need(value['runtimePath'] is None and value['runtimeSha256'] is None and observations['sharedDenied']['warehouse'] is None and value['denials'], '普通主角色没有真实共享拒绝/清缓存')
        for denial in value['denials']:
            need(denial['id'] == 40020 and denial['code'] == 3504 and any(q['id'] == 40020 and q['scope'] == 2 and q['request'] == denial['request'] and q['binding'] == denial['binding'] for q in value['requests']), '普通GUI出现非原Query40020/3504拒绝')
        wrapped = read(path.parent / 'wrapper-result.json'); start = read(path.parent / 'wrapper-start.json')
        need(wrapped['exit_code'] == 0 and wrapped['run_requested'] is True and wrapped['inputs_stable'] is True and wrapped['inputs_before'] == wrapped['inputs_after'] == start['inputs_before'] and wrapped['browser_report_sha256'] == sha(path), '普通浏览器真实子进程包装器未通过')
        current(wrapped['inputs_after'], '普通GUI包装输入')
        need(sha(path.parent / 'run.py') == start['wrapper_sha256'] == contract['normal_wrapper_sha256'] and sha(path.parent / 'tool-at-start.cjs') == contract['tool_sha256'], '普通GUI原执行源未配对')
        for stream in ['stdout', 'stderr']:
            need(sha(path.parent / (stream + '.log')) == wrapped[stream + '_sha256'], '普通GUI原双流日志不同')
    else:
        need(value['denials'] == [], '隔离有效资格出现意外拒绝')
        isolated_parent_evidence(path, value, creator_path)
    return value


def positive(value):
    need(isinstance(value, str) and re.fullmatch(r'[1-9][0-9]*', value) and int(value) <= 2**63 - 1, '数量/金额必须为原正整数64位字符串')
    return int(value)


def same_item(before, after):
    need(all(before[k] == after[k] for k in ['item', 'bound', 'quality', 'sourceKind', 'sourceId', 'locked', 'flags']), '原物品来源/绑定/品质/锁定被重解释')


def validate_fresh_revision(q, op, receipt, before, after):
    """本GUI在Town/Auto关闭且无并行操作，只核新操作，不拿历史回执套新版本。"""
    expected = int(before['revision']) + 1
    need(q['revision'] == before['revision'] and int(op['revision']) == int(receipt['roleRevision']) == int(after['revision']) == expected, 'fresh角色修订不是原before+1')


def validate_item_action(q, receipt, before, after):
    """只比较当前背包+当前已加载页，绝不宣称扫描未加载的整仓。"""
    quantity = positive(q['quantity'])
    need(len(receipt['changes']) == 1 and receipt['changes'][0]['slot'] == q['warehouseSlot'], '原仓位变更不同')
    change = receipt['changes'][0]
    pages = [state['warehouse'] for state in [before, after]]
    need(all(pages[0][k] == pages[1][k] for k in ['scope', 'capacity', 'startSlot', 'limit']), '存取前后不是同一已加载仓页窗口')
    slot = int(q['warehouseSlot']); start = int(pages[0]['startSlot'])
    need(start <= slot < min(start + pages[0]['limit'], int(pages[0]['capacity'])), '目标仓位不在实际加载页')
    slots = [{row['slot']: row['item'] for row in page['items']} for page in pages]
    for state, page, mapped in zip([before, after], pages, slots):
        ids = [row['id'] for row in state['items']] + [row['item']['id'] for row in page['items']]
        need(len(mapped) == len(page['items']) and len(set(ids)) == len(ids), '当前bag+已加载页仓位/UID重复')
    need(slots[0].get(q['warehouseSlot']) == change['before'] and slots[1].get(q['warehouseSlot']) == change['after'], 'receipt前后不等于实际查询页原仓位')
    need({k: v for k, v in slots[0].items() if k != q['warehouseSlot']} == {k: v for k, v in slots[1].items() if k != q['warehouseSlot']}, '非目标当前页原件发生变化')
    original = next((r for r in before['items'] if r['id'] == q['itemInstanceId']), None) if q['action'] == 1 else slots[0].get(q['warehouseSlot'])
    need(original and original['id'] == q['itemInstanceId'] and original['version'] == q['itemVersion'] and quantity <= positive(original['quantity']), '请求不是原完整实例/版本/数量')
    remainder_qty = int(original['quantity']) - quantity
    if q['action'] == 1:
        need(change['before'] is None, '存入目标仓位原先不是空位')
        transferred = change['after']
        remainder = next((r for r in after['items'] if r['id'] == original['id']), None)
        need(before['wallet'] == after['wallet'] and any(r['slot'] == q['warehouseSlot'] and r['item'] == transferred for r in after['warehouse']['items']), '存入实际查询没有原转移事实')
    else:
        remainder = change['after']
        if remainder_qty:
            new_rows = [r for r in after['items'] if r['id'] not in {i['id'] for i in before['items']}]
            need(len(new_rows) == 1, '部分取出没有唯一新UID')
            transferred = new_rows[0]
        else:
            transferred = next((r for r in after['items'] if r['id'] == original['id']), None)
        need(before['wallet'] == after['wallet'] and transferred is not None, '取出没有实际背包原物')
    same_item(original, transferred)
    need(int(transferred['quantity']) == quantity, '转移数量不同')
    if remainder_qty:
        need(remainder and remainder['id'] == original['id'] and int(remainder['quantity']) == remainder_qty and int(remainder['version']) == int(original['version']) + 1, '原余件UID/数量/版本不守恒')
        same_item(original, remainder)
        need(transferred['id'] != original['id'] and transferred['version'] == '1', '拆分新件没有新UID/版本1')
    else:
        need(remainder is None and transferred['id'] == original['id'] and int(transferred['version']) == int(original['version']) + 1, '整件迁移改变UID或没有原版本+1')
    expected_bag = [r for r in before['items'] if q['action'] != 1 or r['id'] != original['id']]
    if q['action'] == 1 and remainder is not None:
        expected_bag.append(remainder)
    if q['action'] == 2:
        expected_bag.append(transferred)
    need(sorted(expected_bag, key=lambda r: r['id']) == sorted(after['items'], key=lambda r: r['id']), '未选择背包原件变化或发生隐式合并')


def isolated_parent_evidence(path, browser, creator_path):
    runtime_path = local(browser['runtimePath']); runtime = read(runtime_path)
    parent_path = runtime_path.parent / 'report.json'; parent = read(parent_path)
    need(sha(runtime_path) == browser['runtimeSha256'] and runtime_path.parent == path.parent.parent and parent['browser_report'] == str(path) and parent['browser_report_sha256'] == sha(path), '隔离原父子报告不配对')
    need(parent['passed'] is True and parent['qa_stopped'] is True and parent['static_stopped'] is True and parent['database_retained'] is True and parent['qa_database_dropped'] is False and parent['main_stopped'] is False and not parent.get('cleanup_errors') and not parent.get('failure'), '隔离父终态/4端口关闭/保库未通过')
    need(runtime['project_root'] == str(ROOT) and runtime['purpose'] == parent['purpose'] == 'isolated_warehouses_gui' and re.fullmatch(r'[0-9a-f]{32}', runtime['run_id']) and runtime['database'] == parent['database'] == 'minimir_dev_game_test_whgui_' + runtime['run_id'] and len(runtime['database'].encode()) < 63, '隔离随机库/marker归属不同')
    need(runtime['ports'] == {'account': 18110, 'game': 18111, 'gm': 18112} and runtime['static_port'] == 18082 and runtime['fixture_sha256'] == sha(ROOT / FIXTURE), '隔离路由或fixture不同')
    need(local(runtime['creator_report']) == creator_path and runtime['creator_report_sha256'] == sha(creator_path) and runtime['artifact_manifest_sha256'] == read(creator_path)['artifact_manifest_sha256'], '隔离不是同一真实候选')
    need(all(parent[k] is True for k in ['main_preserved', 'mongo_preserved', 'source_stable', 'tools_stable', 'go_source_stable', 'artifact_stable']) and parent['main_before'] == parent['main_after'] and parent['mongo_before'] == parent['mongo_after'], '隔离过程主原PID/3实例或Mongo变化')
    for prefix in ['protected', 'tools']:
        need(parent[prefix + '_before'] == parent[prefix + '_after'], '隔离父输入前后漂移')
        current(parent[prefix + '_after'], '隔离父输入')
    need(parent['go_source_before'] == parent['go_source_after'] == runtime['go_source_sha256'] == snapshot('go'), '隔离不是最终Go206原源码')
    q = browser['qualification']; result = q['result']
    need(sha(local(q['helperPath'])) == q['helperSha256'] == QA_RUNNER_SHA and q['exitCode'] == 0 and result['passed'] is True, '隔离资格原helper/终态不符')
    for stream in ['stdout', 'stderr']:
        need(sha(path.parent / ('qualify.' + stream + '.log')) == q[stream + 'Sha256'], '资格原双流SHA不同')
    need(json_output(path.parent / 'qualify.stdout.log') == result and result['role_id'] == browser['createdRoles'][0]['roleId'] and result['source'] == 'development_fixture' and result['warehouse_results_seeded'] is False and result['natural_supply_or_real_payment_proven'] is False and result['credentials_written'] is False, '资格结果不是正常创角前态或越界业务播种')
    proof_path = local(result['qualification']); need(proof_path.parent == runtime_path.parent / 'qualifications' and sha(proof_path) == result['qualification_sha256'], '隔离资格before/after原件不配对')
    proof = read(proof_path); before, after = proof['before'], proof['after']
    need(proof['_id'] == result['role_id'] and proof['run_id'] == runtime['run_id'] and proof['purpose'] == runtime['purpose'] and proof['fixture_sha256'] == sha(ROOT / FIXTURE), '资格角色/运行/fixture身份不同')
    expected = copy.deepcopy(before)
    for key in ['privilege_until', 'wallet', 'items']:
        expected[key] = after[key]
    need(expected == after and before['revision'] == after['revision'] == 1 and before['economic_command_seq'] == after['economic_command_seq'] == 0 and before['source_seq'] == after['source_seq'] == 0 and before['owner_instance'] == after['owner_instance'] == '', '资格覆盖原角色其它字段/已入场/已操作')
    need(after['items'][:-2] == before['items'] and [(r['item_id'], r['quantity'], r['bound'], r['version'], r['source_kind']) for r in after['items'][-2:]] == [(218, 1, False, 1, 'development_fixture'), (900001, 5, False, 1, 'development_fixture')], '资格改变原初始物或捏造不同测试供给')
    cost = read(ROOT / FIXTURE)['warehouses']['personal_expansion_yuanbao_units']
    need(after['wallet']['gold'] == result['existing_gold'] == 1000 and after['wallet']['yuanbao_units'] == result['existing_yuanbao_units'] == cost + 1000 and all(before['wallet'][k] == after['wallet'][k] for k in ['bound_gold', 'bound_yuanbao_units']), '已有测试资金或绑定钱包合同不同')


class Archive:
    def __init__(self):
        self.jobs = {}
        self.sources = set()

    def add(self, source, target, scope):
        source, target = local(source), Path(target)
        need(source.is_file() and source.suffix.lower() not in {'.exe', '.dll', '.lib', '.wil', '.wix', '.wzl', '.wzx', '.pak', '.zip', '.7z', '.rar', '.ttf', '.ttc', '.wav', '.mp3', '.bin'}, '拒归档构建/引擎/字体/二进制包')
        need(not target.is_absolute() and '..' not in target.parts and not set(target.parts) & {'node_modules', 'library', 'temp', 'build', 'candidate', 'client-project', 'previous-build', 'failed-build'}, '归档目标包含缓存/整个产物目录')
        name = relative(source)
        if name in self.sources:
            return
        compressed = source.suffix == '.jsonl' or source.suffix in {'.txt', '.log', '.json'} and source.stat().st_size >= 131072
        if compressed:
            target = target.with_name(target.name + '.gz')
        need(target.as_posix() not in self.jobs, '归档目标重复')
        self.jobs[target.as_posix()] = {'source': name, 'target': target.as_posix(), 'scope': scope, 'encoding': 'gzip' if compressed else 'identity', 'raw_sha256': sha(source), 'raw_bytes': source.stat().st_size}
        self.sources.add(name)

    def bundle(self, report, target, scope):
        report = local(report)
        self.add(report, Path(target) / report.name, scope)
        value = read(report)
        names = ['wrapper.json', 'started.json', 'before.json', 'events.jsonl', 'creator.log', 'candidate-files.json', 'creator-build-candidate.json',
                 'served-build-before.json', 'served-build-after.json', 'visual-review.json', 'handoff.json', 'run.py', 'tool-at-start.cjs',
                 'loaded-public.json', 'stdout.txt', 'stderr.txt', 'stdout.log', 'stderr.log', 'actual.log', 'command.json', 'mongo-status.json',
                 'mongo-status.stdout.json', 'mongo-status.stderr.txt', 'qualify.stdout.log', 'qualify.stderr.log', 'failure.png']
        for name in names:
            if (report.parent / name).is_file():
                self.add(report.parent / name, Path(target) / name, scope)
        for name in value.get('screenshots', []):
            need(Path(name).name == name and name.endswith('.png'), '历史截图名称必须为同目录PNG')
            self.add(report.parent / name, Path(target) / 'screenshots' / name, scope)
        for parent in [value, read(report.parent / 'wrapper.json') if (report.parent / 'wrapper.json').is_file() else {}]:
            for row in parent.get('commands', []):
                for stream in ['stdout', 'stderr']:
                    if stream in row:
                        p = report_log(row[stream], report)
                        if stream + '_sha256' in row:
                            need(sha(p) == row[stream + '_sha256'], '支持原日志SHA不符')
                        else:
                            need(p.parent == report.parent, '无producer摘要的支持日志必须是该报告同目录原件')
                        self.add(p, Path(target) / p.name, scope)
            if 'events_sha256' in parent:
                need(sha(report.parent / 'events.jsonl') == parent['events_sha256'], '历史events SHA不符')

    def apply(self):
        OUT.mkdir(parents=True, exist_ok=False)
        records = []
        for job in self.jobs.values():
            source, target = ROOT / job['source'], OUT / job['target']
            need(sha(source) == job['raw_sha256'], '归档前原件改变')
            target.parent.mkdir(parents=True, exist_ok=True)
            if job['encoding'] == 'gzip':
                with source.open('rb') as src, target.open('xb') as dst, gzip.GzipFile(filename='', mode='wb', fileobj=dst, mtime=0) as stream:
                    shutil.copyfileobj(src, stream)
                need(digest(gzip.decompress(target.read_bytes())) == job['raw_sha256'], 'gzip原字节不等')
            else:
                shutil.copyfile(source, target)
                need(sha(target) == job['raw_sha256'], '复制原字节不等')
            records.append({**job, 'file_sha256': sha(target), 'file_bytes': target.stat().st_size})
        return records


def curate(args, protocol_files):
    archive = Archive()
    selected = {'go': args.go_report, 'client': args.client_report, 'creator': args.creator_report, 'publication': args.publication_report,
                'normal-gui': args.normal_gui_report, 'isolated-gui': args.isolated_gui_report}
    for kind, path in selected.items():
        archive.bundle(path, 'current/' + kind, 'current_' + kind + '_selected_terminal')
    archive.add(args.gui_contract, 'support/gui-producer-contract.json', 'explicit_selected_frozen_gui_contract_not_execution')
    archive.bundle(read(args.gui_contract)['static_report'], 'support/gui-tool-static-current', 'current_gui_tool_pure_not_browser_business')
    for mode, path in [('normal', args.normal_gui_report), ('isolated', args.isolated_gui_report)]:
        value = read(path)
        for name in value['screenshots']:
            archive.add(path.parent / name, 'current/' + mode + '-gui/screenshots/' + name, 'actual_touch_ui_screenshot_reviewed')
        for name in ['wrapper-start.json', 'wrapper-result.json']:
            p = path.parent / name
            if p.is_file():
                archive.add(p, 'current/' + mode + '-gui/' + name, 'actual_gui_child_process_metadata')
    isolated_browser = read(args.isolated_gui_report)
    runtime_path = local(isolated_browser['runtimePath'])
    parent_dir = runtime_path.parent
    archive.bundle(parent_dir / 'report.json', 'current/isolated-parent', 'actual_isolated_parent_main_preserved_qa_stopped_database_retained')
    for name in ['runtime.json', 'static-runtime.json', 'build.log', 'node-path.log', 'browser.log', 'server.log', 'static.log']:
        archive.add(parent_dir / name, 'current/isolated-parent/' + name, 'isolated_runtime_support_no_credentials_or_database_dump')
    proof_path = local(isolated_browser['qualification']['result']['qualification'])
    archive.add(proof_path, 'current/isolated-parent/qualification.json', 'actual_gui_qualification_before_after_not_natural_supply_or_real_payment')
    for name in ['runtime-before.json', 'runtime-after.json']:
        archive.add(WORK / name, 'current/publication/' + name, 'current_cutover_support')
    for path in protocol_files:
        archive.add(path, 'current/client/protocol/' + path.name, 'codec_only_not_business')
    archive.bundle(WORK / 'configuration-run-01/report.json', 'configuration/run-01', 'original_configuration_140_python_and_11_commands')
    for name in ['configuration-baseline.json', 'configuration-edit-plan.json', 'configuration-editor-report.json', 'configuration-review.json', 'current-go-test-adaptations.json',
                 'client/production-freeze.json', 'client-peer-review-pre-freeze.json', 'client-peer-review-final.json', 'qa-runner-handoff-final.json']:
        archive.add(WORK / name, 'support/' + name, 'support_original_scope_not_full_business')
    plan = read(WORK / 'curate-plan.json')
    for source in plan['private_sources']:
        archive.add(ROOT / source, 'sources/' + source, 'reproducible_private_tool_source')
    for path in read(WORK / 'source-allowlist.json')['allowed_paths']:
        p = ROOT / path
        if p.is_file() and p.suffix in {'.py', '.cjs', '.mjs', '.json'} and path.startswith('tools/'):
            archive.add(p, 'sources/' + path, 'reproducible_checked_tool_source')
    for entry in plan['history_reports']:
        p = ROOT / entry['path']
        if p.is_file():
            archive.bundle(p, 'history/' + p.parent.relative_to(WORK).as_posix(), entry['scope'])
            if (p.parent / 'runtime.json').is_file() and read(p).get('purpose') == 'isolated_warehouses_gui':
                for name in ['runtime.json', 'static-runtime.json', 'build.log', 'node-path.log', 'browser.log', 'server.log', 'static.log']:
                    q = p.parent / name
                    if q.is_file():
                        archive.add(q, 'history/' + q.relative_to(WORK).as_posix(), entry['scope'])
                child = p.parent / 'browser/report.json'
                if child.is_file() and 'qualification' in read(child):
                    proof = local(read(child)['qualification']['result']['qualification'])
                    archive.add(proof, 'history/' + proof.relative_to(WORK).as_posix(), 'historical_isolated_qualification_only_not_business_pass')
    for entry in plan['history_files']:
        p = ROOT / entry['path']
        need(p.is_file(), '固定历史原件缺失：' + str(p))
        archive.add(p, 'history/' + p.relative_to(WORK).as_posix(), entry['scope'])
    for entry in plan['support_reports']:
        p = ROOT / entry['path']
        need(p.is_file(), '固定支持报告缺失：' + str(p))
        archive.bundle(p, 'support/' + p.parent.relative_to(WORK).as_posix(), entry['scope'])
    for entry in plan['support_files']:
        p = ROOT / entry['path']
        archive.add(p, 'support/' + p.relative_to(WORK).as_posix(), entry['scope'])
    for path in args.support_report:
        p = local(path); need(p.is_relative_to(WORK), '支持报告仅本阶段')
        archive.bundle(p, 'support/selected/' + p.parent.relative_to(WORK).as_posix(), 'explicit_support_not_full_current')
    for path in args.history_report:
        p = local(path); need(p.is_relative_to(WORK), '历史报告仅本阶段')
        archive.bundle(p, 'history/selected/' + p.parent.relative_to(WORK).as_posix(), 'explicit_historical_original_not_reclassified_pass')
    return archive


def evidence_attributes():
    raw = git('show', BASE + ':.gitattributes')
    need(digest(raw) == ATTR_BASE and raw.endswith(b'\n') and ATTR_APPEND not in raw, 'Git证据BASE属性不同')
    need((ROOT / '.gitattributes').read_bytes() == raw + ATTR_APPEND, '只能追加本阶段精确-text行，不改全局Git或日志')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(); mode.add_argument('--check', action='store_true'); mode.add_argument('--apply', action='store_true')
    parser.add_argument('--go-report', type=local, default=WORK / 'go/full-2/report.json')
    parser.add_argument('--client-report', type=local, default=WORK / 'client/full-3/report.json')
    parser.add_argument('--creator-report', type=local, default=WORK / 'creator/full-3/report.json')
    parser.add_argument('--publication-report', type=local, default=WORK / 'runtime-control/report.json')
    parser.add_argument('--publication-source-review', type=local, default=WORK / 'cutover-selector-review-3.json')
    parser.add_argument('--normal-gui-report', type=local, required=True)
    parser.add_argument('--isolated-gui-report', type=local, required=True)
    parser.add_argument('--gui-contract', type=local, default=WORK / 'browser/report-contract-3.json')
    parser.add_argument('--gui-contract-sha256', required=True)
    parser.add_argument('--support-report', type=local, action='append', default=[])
    parser.add_argument('--history-report', type=local, action='append', default=[])
    args = parser.parse_args()
    need(git('rev-parse', 'HEAD').strip().decode() == BASE, 'HEAD不是已推送rule22基线')
    evidence_attributes()
    config = configuration_evidence()
    go, go_stats = go_evidence(args.go_report)
    client, client_stats, protocol_files = client_evidence(args.client_report)
    delta = protocol_delta()
    creator, layout = creator_evidence(args.creator_report)
    selected = {}
    for group, path in [('go', args.go_report), ('client', args.client_report), ('creator', args.creator_report)]:
        need(path.name == 'report.json' and path.parent.parent == WORK / group and re.fullmatch(r'full-[1-9][0-9]*', path.parent.name), 'CLI selector必须是本阶段该检查组精确full目录')
        selected[group] = path.parent.name.removeprefix('full-')
    publication = publication_evidence(args.publication_report, creator, args.publication_source_review, selected)
    normal = gui_evidence(args.normal_gui_report, 'normal', args.creator_report, args.gui_contract, args.gui_contract_sha256)
    isolated = gui_evidence(args.isolated_gui_report, 'isolated', args.creator_report, args.gui_contract, args.gui_contract_sha256)
    archive = curate(args, protocol_files)
    selectors = {key: {'path': relative(path), 'sha256': sha(path)} for key, path in {
        'go_report': args.go_report, 'go_events': args.go_report.parent / 'events.jsonl', 'client_report': args.client_report,
        'creator_report': args.creator_report, 'publication_report': args.publication_report, 'normal_gui_report': args.normal_gui_report,
        'isolated_gui_report': args.isolated_gui_report, 'normal_gui_wrapper': args.normal_gui_report.parent / 'wrapper-result.json',
        'normal_loaded_public': args.normal_gui_report.parent / 'loaded-public.json',
        'isolated_loaded_public': args.isolated_gui_report.parent / 'loaded-public.json',
        'isolated_gui_parent': local(isolated['runtimePath']).parent / 'report.json',
        'isolated_runtime': local(isolated['runtimePath']), 'configuration_report': WORK / 'configuration-run-01/report.json',
        'configuration_review': WORK / 'configuration-review.json', 'gui_contract': args.gui_contract,
        'publication_source_review': args.publication_source_review}.items()}
    now = dt.datetime.now(dt.timezone.utc)
    stage = {'schema_version': 1, 'date': now.astimezone(dt.timezone(dt.timedelta(hours=8))).date().isoformat(), 'recorded_at_utc': now.isoformat(),
             'status': 'limited_passed', 'full_p0_completed': False, 'production_ready': False, 'runtime_ready': False,
             'scope': '双仓储与五角色并发的限定开发组合；正常免费个人仓和隔离共享/扩页资格分别验收，不冒充权益购买、真实支付、运营合服、真机或容量',
             'runtime': {'rule_version': 23, 'preview': config['preview'], 'fixture_sha256': config['fixture_sha256'], 'public_sha256': config['public_sha256'],
                         'skills': 22, 'bootstrap_mode': 'legacy_untrained_pet', 'casting_supply_enabled': False, 'summon17_enabled': False,
                         'table_count': read(ROOT / 'config/reports/配置校验报告.json')['table_count'],
                         'record_count': read(ROOT / 'config/reports/配置校验报告.json')['record_count'],
                         'release_blockers': len(read(ROOT / 'config/reports/配置校验报告.json')['release_blockers']),
                         'entitlement_purchase_verified': False, 'actual_payment_verified': False, 'operational_merge_verified': False,
                         'physical_mobile_verified': False, 'capacity_verified': False},
             'verification': {'go_full_mongo_ws_race_passed_including_subtests': go_stats['passed'], 'go_failed': 0, 'go_skipped': 0, 'go_inputs': go_stats['inputs'],
                              'configuration_python_passed': config['python'], 'configuration_inputs': 27, 'javascript_ui_app_passed': client_stats['app'],
                              'javascript_scene_passed': client_stats['scene'], 'javascript_total_passed': client_stats['total'], 'client_inputs': client_stats['inputs'],
                              'client_external_engine_inputs': client_stats['engine_inputs'], 'client_external_font_inputs': client_stats['font_inputs'], 'pb_messages': client_stats['pb_messages'],
                              'pb_long_fields': client_stats['pb_long_fields'], 'pb_boundary_cases': client_stats['pb_boundary_cases'], 'pb_protocol_ids': client_stats['pb_protocol_ids'],
                              'creator_seconds': creator['seconds'], 'creator_exit_code': 36, 'creator_time_scope': '实际Creator进程加构建后SHA验证；wrapper另含构建前旧包SHA，不是纯编译时长',
                              'saved_layouts_unchanged_against_base': 9, 'saved_layouts_intentionally_changed': 1, 'gameplay_excel_changed_cells': 2,
                              'formulas_preserved': 460, 'unchanged_workbook_zip_members': 102, 'normal_gui_checks': len(normal['checks']), 'isolated_gui_checks': len(isolated['checks'])},
             'layout': layout, 'protocol_delta': delta, 'selectors': selectors,
             'gui': {'normal': {'checks': normal['checks'], 'errors': normal['errors'], 'warnings': normal['warnings'], 'screenshots': normal['screenshots']},
                     'isolated': {'checks': isolated['checks'], 'errors': isolated['errors'], 'warnings': isolated['warnings'], 'screenshots': isolated['screenshots'],
                                  'qualification_scope': '随机库正常创角后、首次Enter前的明确已有权益/资金/2测试物；真实before/after单列，不是自然供给或真实支付'}},
             'scope_exclusions': ['权益购买', '真实付费', '运营合服', '真机', '容量', '完整P0', '自然非绑元宝供给'],
             'qualification_evidence_boundary': '原WS49资格仅对应源码和raw断言，未保存逐UID资格原JSON；隔离GUI资格则保存独立实际before/after，不混同范围',
             'known_visual_limitations': ['原MapTravel确认Message的单长行可左右截断，本轮保持旧源并单列限制；不宣称所有界面无截断或无问题'],
             'history_note': '原失败、未观察guard的旧回归、组件注入防御与入口可达性纠正均保留原scope，不追认为已复现生产红绿', 'evidence_sha256': {}}
    if not args.apply:
        print(json.dumps({'passed': True, 'check_only': True, 'planned_files': len(archive.jobs), 'selectors': selectors, 'verification': stage['verification']}, ensure_ascii=False))
        return
    need(not OUT.exists() and not FORMAL.exists(), '拒绝覆盖已有正式报告/归档')
    rows = archive.apply()
    manifest = {'schema_version': 1, 'baseline': BASE, 'files': rows}
    paths = {r['source']: {'path': relative(OUT / r['target']), 'encoding': r['encoding'], 'raw_sha256': r['raw_sha256'], 'file_sha256': r['file_sha256']} for r in rows}
    for name, value in [('evidence-manifest.json', manifest), ('evidence-paths.json', paths)]:
        (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8', newline='\n')
    (OUT / 'README.md').write_text('本目录仅保存本阶段限定验收的原始报告、日志、截图与必要工具源。大文本采用无损gzip，manifest同时记录原字节和包装SHA。\n\n共享权益、已有资金及两件测试物为隔离资格；不认证权益购买、真实付费、自然供给、运营合服、真机、容量或完整P0。主保持22技能、legacy免费未训宠、护符及17关闭。\n', encoding='utf-8', newline='\n')
    stage['evidence_sha256'] = tree(OUT)
    FORMAL.write_text(json.dumps(stage, ensure_ascii=False, indent=2) + '\n', encoding='utf-8', newline='\n')
    print(json.dumps({'passed': True, 'applied': True, 'formal': relative(FORMAL), 'evidence_files': len(stage['evidence_sha256'])}, ensure_ascii=False))


if __name__ == '__main__':
    try:
        main()
    except (ValueError, KeyError, FileNotFoundError, AssertionError) as error:
        print(json.dumps({'passed': False, 'pending_or_invalid': True, 'error': str(error)}, ensure_ascii=False))
        raise SystemExit(1)
